import openpyxl
import requests

# Read the data from the Excel workbook
workbook = openpyxl.load_workbook('C:\\Users\\anyone\\Downloads\\updated names.xlsx')
sheet = workbook.active

# Read the values from the first row of the sheet
values = []
for cell in sheet[1]:
    values.append(cell.value)

# Define the message with placeholders for the values
message = "The values are: {value1}, {value2}, {value3}"

# Replace the placeholders with the actual values
formatted_message = message.format(value1=values[1], value2=values[2], value3=values[3])
receipt_address = values[0]

# Send the message using the MNOTIFY API
url = "https://api.mnotify.net/smsapi"
payload = {
    "key": "pX88BJbnvGDH5PlAuVi70IuFE",
    "to": receipt_address,
    "msg": formatted_message,
    "sender_id": "IHP23"
}
response = requests.post(url, data=payload, verify=False)

print(response.status_code)
